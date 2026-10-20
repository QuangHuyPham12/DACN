"""Build the factual report and diagram assets from the implemented SOLE version."""
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/defense'
BUILD = ROOT / 'storage/app/defense-build'
OUT.mkdir(parents=True, exist_ok=True)
BUILD.mkdir(parents=True, exist_ok=True)
FONT = 'C:/Windows/Fonts/arial.ttf'
BOLD = 'C:/Windows/Fonts/arialbd.ttf'
INK, GREEN, ORANGE, BG = '#22271f', '#324e3f', '#c46c3d', '#f7f5f0'

def diagram(name, title, boxes, arrows=(), footer=''):
    im=Image.new('RGB',(1600,850),'white'); d=ImageDraw.Draw(im)
    d.text((55,35),title,font=ImageFont.truetype(BOLD,38),fill=INK)
    for a,b in arrows:
        d.line([a,b],fill=GREEN,width=4)
        x,y=b; dx,dy=b[0]-a[0],b[1]-a[1]
        if abs(dx)>abs(dy):
            s=1 if dx>0 else -1; d.polygon([(x,y),(x-16*s,y-9),(x-16*s,y+9)],fill=GREEN)
        else:
            s=1 if dy>0 else -1; d.polygon([(x,y),(x-9,y-16*s),(x+9,y-16*s)],fill=GREEN)
    for rect,text,dark in boxes:
        d.rounded_rectangle(rect,18,fill=GREEN if dark else BG,outline='#c4cabe',width=2)
        lines=text.split('\n'); f=ImageFont.truetype(BOLD if dark else FONT,29)
        for i,line in enumerate(lines):
            w=d.textlength(line,font=f)
            d.text(((rect[0]+rect[2]-w)/2,(rect[1]+rect[3])/2-len(lines)*20+i*42),line,font=f,fill='white' if dark else INK)
    d.text((55,790),footer,font=ImageFont.truetype(FONT,23),fill='#596257')
    path=OUT/(name+'.png'); im.save(path); return path

diagram('architecture','Kiến trúc của website SOLE',[
    ((60,170,490,350),'Trình duyệt\nBlade HTML CSS JavaScript',False),
    ((590,170,1030,350),'PHP Laravel\nRoute và controller',True),
    ((590,440,1030,630),'Service nghiệp vụ\nTransaction và khóa dữ liệu',True),
    ((1140,440,1540,630),'MySQL 8\n34 bảng dữ liệu',False),
    ((60,440,490,630),'Tác vụ mỗi phút\nHết hạn QR và email',False)],
    [((490,260),(590,260)),((810,350),(810,440)),((1030,535),(1140,535)),((490,535),(590,535))],
    'Hai guard đăng nhập riêng cho khách và nhân viên; kiểm tra quyền ở máy chủ.')
diagram('use-cases','Vai trò và nhóm chức năng',[
    ((50,150,340,300),'Khách hàng',True),((50,370,340,520),'Nhân viên',True),((50,590,340,740),'Quản trị viên',True),
    ((500,120,1530,310),'Xem lọc giày · giỏ hàng · COD và QR\nTài khoản · đánh giá đã mua · yêu cầu đổi trả',False),
    ((500,350,1530,540),'POS · xử lý đơn · vị trí giao hàng\nKhách hàng · tiếp nhận yêu cầu đổi trả',False),
    ((500,580,1530,755),'Sản phẩm và kho · nhập và trả hàng · phân quyền\nBáo cáo · kiểm duyệt đánh giá · quản lý email',False)],
    [((340,225),(500,225)),((340,445),(500,445)),((340,665),(500,665))])
diagram('payment-flow','Đặt hàng QR và xử lý hết hạn',[
    ((60,160,470,310),'Khách đặt đơn QR\nPHP tính tổng tiền',False),
    ((580,160,1010,310),'Lưu đơn và giữ hàng\nHạn thanh toán 30 phút',True),
    ((1120,160,1540,310),'Khách báo đã chuyển\nNhân viên đối soát',False),
    ((580,460,1010,650),'Đến hạn và chưa xử lý\nTự hủy và trả phần giữ',False),
    ((1120,460,1540,650),'Đã nhận đủ tiền\nTiếp tục xử lý giao hàng',True)],
    [((470,235),(580,235)),((1010,235),(1120,235)),((795,310),(795,460)),((1330,310),(1330,460))],
    'Không tự hủy đơn đã báo chuyển, đã thu hoặc đã xác nhận. Tiền đến muộn chờ hoàn trên đơn hủy.')
diagram('data-relations','Các quan hệ dữ liệu chính',[
    ((50,170,410,330),'products\nproduct_variants',True),((610,170,980,330),'sales_orders\nsales_order_items',True),
    ((1180,170,1550,330),'payments\nsales_returns',False),((50,490,410,650),'inventory_movements\npurchase_orders và items',False),
    ((610,490,980,650),'customers\ncustomer_accounts',True),((1180,490,1550,650),'reviews · support_requests\norder_messages',False)],
    [((410,250),(610,250)),((980,250),(1180,250)),((230,330),(230,490)),((795,490),(795,330)),((980,570),(1180,570)),((1365,330),(1365,490))],
    'Sơ đồ rút gọn; ERD đầy đủ và ràng buộc trong docs/DATABASE.md và migration.')

# Reuse actual screenshots; no invented interface or customer feedback.
for source,target,height in [('support-request.jpg','support-screen.png',1080),('password-recovery.jpg','password-screen.png',900),('review-moderation.jpg','review-screen.png',950),('storefront-desktop.jpg','storefront-screen.png',900),('delivery-tracking.jpg','tracking-screen.png',1100)]:
    p=ROOT/'docs/screenshots'/source
    if p.exists():
        im=Image.open(p).convert('RGB'); im.crop((0,0,im.width,min(height,im.height))).save(OUT/target)

pages=[
('Mở đầu và phạm vi',[
('p','Nhóm xây dựng SOLE để quản lý một cửa hàng giày và một kho, đồng thời cho khách mua trực tuyến. Hệ thống quản lý theo SKU gồm mẫu giày, size và màu; mỗi giao dịch nhập, bán, điều chỉnh và trả hàng có chứng từ để đối chiếu.'),
('p','Phần mềm hiện chạy bằng PHP Laravel, JavaScript và MySQL. Các chức năng mua hàng, giữ kho, thanh toán và giao hàng đã được triển khai cùng kiểm thử tự động. Dữ liệu, ảnh minh họa và các giao dịch DEMO phục vụ đồ án; không đại diện hoạt động kinh doanh thực tế.'),
('table',[['Vai trò','Phạm vi'],['Khách hàng','Xem và lọc giày, giỏ hàng, đặt COD/QR, tài khoản, đánh giá đã mua, yêu cầu đổi/trả'],['Nhân viên','POS, xử lý đơn và vận chuyển, tra cứu tồn, khách hàng, phản hồi yêu cầu'],['Quản trị viên','Danh mục, nhập/điều chỉnh kho, hoàn tiền, phân quyền, báo cáo và kiểm duyệt']]),
('p','Báo cáo mô tả bản triển khai ngày 21/10/2026, đối chiếu đề cương hiện có của môn Đồ án chuyên ngành Công nghệ phần mềm. Thông tin hành chính, mã sinh viên và nhận xét giảng viên cần được nhóm bổ sung theo hồ sơ chính thức; báo cáo không tự điền điểm hoặc xác nhận.')]),
('Chương 1 Cơ sở công nghệ',[
('p','Laravel tổ chức ứng dụng theo MVC. Route và middleware xác định điểm truy cập, phiên đăng nhập và quyền; controller xác thực dữ liệu; service tập trung các thay đổi kho, đơn và tiền trong transaction; Eloquent ánh xạ dữ liệu quan hệ. Blade tạo HTML, JavaScript tăng tương tác trên trình duyệt.'),
('img','architecture.png'),
('table',[['Thành phần','Bản triển khai'],['Máy chủ','PHP 8.3.35 và Laravel 12'],['Trình duyệt','Blade, HTML, CSS, JavaScript thuần'],['Cơ sở dữ liệu','MySQL 8.0.46 tại 127.0.0.1:3309'],['Kiểm thử','PHPUnit trên SQLite và MySQL riêng']]),
('p','MySQL là database của website. SQLite chỉ dùng trong một số bài kiểm thử và giữ bản nguồn trước khi nâng cấp. MySQL Workbench là công cụ kết nối để xem và quản trị dữ liệu.')]),
('Chương 2 Yêu cầu và phân quyền',[
('img','use-cases.png'),
('p','Khách và nhân viên dùng hai guard riêng. Đăng nhập khách không cấp quyền quản trị. Nhân viên không sửa danh mục, không điều chỉnh kho, không duyệt đánh giá hoặc xuất báo cáo. Admin có các quyền này nhưng không tự khóa tài khoản đang dùng.'),
('p','Mọi giá bán, giảm giá, phí giao, quyền và lượng hàng đều được kiểm tra tại server. Hệ thống phải tránh bán âm, trừ kho hai lần và ghi nhận tiền hai lần khi người dùng gửi lại thao tác. Giao diện có nhãn biểu mẫu, phản hồi lỗi và bố cục dùng được trên điện thoại.')]),
('Thiết kế cơ sở dữ liệu',[
('img','data-relations.png'),
('p','Database sole_store có 34 bảng: 25 bảng nghiệp vụ và 9 bảng hạ tầng. Nhóm bảng chính gồm tài khoản, danh mục, sản phẩm/biến thể/ảnh, khách/nhà cung cấp, nhập/bán, thanh toán, sổ kho, ưu đãi, trả hàng, đánh giá, yêu cầu hỗ trợ và email đơn hàng.'),
('p','SKU và tổ hợp sản phẩm–size–màu là duy nhất. Mỗi đơn một khoản thu; mã giao dịch ngân hàng không dùng cho hai đơn. Một tài khoản đánh giá mỗi sản phẩm một lần; mỗi đơn một yêu cầu hỗ trợ và một email xác nhận theo loại. Khóa ngoại bảo vệ lịch sử chứng từ.'),
('p','Tiền lưu DECIMAL(15,0), xử lý bằng số nguyên VND. Chi tiết đơn lưu tên, SKU, size/màu và giá tại lúc bán. Sửa sản phẩm không làm đổi đơn cũ. Token khôi phục khách được băm và lưu riêng token nhân viên. Chi tiết cột và ERD đầy đủ nằm trong docs/DATABASE.md và migration.')]),
('Quy trình đặt hàng và thanh toán',[
('img','payment-flow.png'),
('p','Khách chọn đúng biến thể, nhập thông tin nhận, áp ưu đãi và chọn COD hoặc QR. PHP tính tổng, lưu đơn cùng snapshot. QR chỉ bật khi đã có tài khoản ngân hàng hợp lệ. Nút báo chuyển khoản không tự tạo payment; nhân viên đối chiếu sao kê, số tiền và mã giao dịch trước khi xác nhận.'),
('p','Đơn QR mới chưa báo chuyển, chưa thu và chưa được xác nhận có hạn mặc định 30 phút. Scheduler chạy mỗi phút; thao tác xem đơn và báo chuyển cũng kiểm tra hạn. Hủy giải phóng phần giữ kho và lượt ưu đãi một lần. Nếu tiền đến muộn, giữ đơn hủy và ghi nhận chờ hoàn toàn bộ; không khôi phục đơn hoặc hàng giữ.')]),
('Quản lý kho giao hàng và tài chính',[
('p','Tồn sổ sách gồm hàng chưa xuất bán, kể cả hàng giữ chỗ/đang giao. Số có thể bán bằng tồn sổ sách trừ giữ chỗ. COD giữ khi xác nhận; QR giữ khi tạo. Chỉ hoàn tất mới xuất sổ sách và giải phóng phần giữ, kể cả đã thu QR trước.'),
('table',[['Thời điểm','Tồn sổ sách','Giữ chỗ','Có thể bán'],['Trước đơn','10','0','10'],['Giữ 2 đôi','10','2','8'],['Hoàn tất đơn','8','0','8']]),
('p','Giao hàng chuyển tuần tự: chờ xác nhận, xác nhận, chuẩn bị, đang giao, đã giao. QR phải thu đủ trước khi bàn giao. Giao thất bại có thể thử lại hoặc nhận hàng hoàn về rồi hủy; không trả kho giữ khi hàng vẫn đang đi.'),
('p','Tổng thanh toán = tiền hàng − giảm giá + phí giao. Phí mặc định 30.000đ, miễn từ 1.500.000đ tiền hàng trước giảm. Doanh thu tính theo đơn hoàn tất và trừ phiếu trả. Dòng tiền tính theo ngày thu/hoàn; tiền trả trước chưa phải doanh thu. Báo cáo chưa tính lợi nhuận vì chưa có phương pháp giá vốn.'),
('p','Trả đơn hoàn tất hoàn tiền hàng sau giảm, không hoàn phí giao. Hủy đơn trả trước hoàn toàn bộ số đã thu, gồm phí giao. Cả hai là xác nhận giao dịch thực tế do cửa hàng thực hiện, không tự chuyển tiền ngân hàng.')]),
('Bảo vệ tài khoản và khôi phục mật khẩu',[
('p','Luồng quên mật khẩu trả cùng thông báo cho email có và không có tài khoản. Liên kết tạo từ APP_URL, dùng token ngẫu nhiên được băm, hết hạn 60 phút và chỉ sử dụng một lần. Hàng token được khóa khi đặt lại mật khẩu để tránh hai yêu cầu dùng đồng thời.'),
('img','password-screen.png'),
('p','Đặt lại hoặc đổi mật khẩu tăng auth_version. Middleware kiểm tra phiên trên toàn website, bao gồm checkout; phiên cũ mất hiệu lực ở yêu cầu tiếp theo. Hệ thống còn có CSRF, giới hạn đăng nhập/reset/chat, validation server, escape nội dung và tách dữ liệu theo chủ tài khoản. Thư reset hiện dùng mail log nội bộ, cần SMTP để gửi thật.')]),
('Đánh giá sản phẩm và nội dung danh mục',[
('p','Chỉ tài khoản có đơn hoàn tất chứa sản phẩm mới gửi được đánh giá từ 1 đến 5 sao. Nội dung mới/chỉnh sửa ở trạng thái chờ duyệt; chỉ admin duyệt hoặc ẩn. Điểm trung bình chỉ dùng đánh giá đã duyệt. Tên hiển thị rút gọn và nội dung được escape; không công khai email hay địa chỉ.'),
('img','review-screen.png'),
('p','Admin có thể thêm chất liệu đã xác nhận, form, cách chăm sóc và ghi chú size cho từng mẫu. Thư viện tối đa 8 ảnh JPG/PNG/WebP, mỗi ảnh không quá 4 MB, cho phép đổi bìa và xóa từng ảnh. Ảnh thật phải đúng mẫu, màu và có quyền sử dụng. Chưa có ảnh thật thì giữ nhãn minh họa, không tạo thông số hoặc lời đánh giá giả.')]),
('Yêu cầu đổi size và trả hàng',[
('img','support-screen.png'),
('p','Khách gửi yêu cầu từ đơn của mình trong trang tài khoản. Yêu cầu gồm lý do và size mong muốn nếu đổi; một yêu cầu trên mỗi đơn. Cửa hàng tiếp nhận, trao đổi, đồng ý hoặc từ chối. Phản hồi gần nhất được hiển thị cho đúng khách sở hữu đơn.'),
('p','Gửi yêu cầu hoặc đồng ý tiếp nhận chưa thay đổi kho/tiền. Admin phải nhận hàng, kiểm tra và lập phiếu trả trước khi đóng yêu cầu. Bản hiện tại trả toàn bộ đơn, đổi size bằng trả đơn và lập đơn mua mới; chưa trả từng dòng hoặc tự bù chênh lệch. Ảnh chụp trang này là dữ liệu DEMO ghi rõ trong nội dung.')]),
('Chương 3 Cài đặt và tác vụ nền',[
('p','Chạy start.ps1 trong shoe-store để khởi động MySQL local nếu cần, cập nhật migration an toàn, bật scheduler và website. Trên máy hiện tại, database sole_store ở cổng 3309; sole_store_test được dành cho test, không trộn với database website.'),
('table',[['Việc cần làm','Lệnh hoặc cấu hình'],['Khởi động','PowerShell: .\\start.ps1'],['Chỉ tác vụ nền','.\\start-worker.ps1 hoặc php artisan schedule:work'],['Kiểm tra lịch','php artisan schedule:list'],['Áp dụng cấu trúc','php artisan migrate --force'],['Thử hết hạn QR','php artisan store:expire-qr'],['Thử hàng đợi email','php artisan store:send-order-mail']]),
('p','Email xác nhận được ghi vào outbox trong transaction đặt hàng, sau đó tác vụ gửi và thử lại tối đa ba lần nếu lỗi. Chế độ log/array tạo bản xem thử, không báo gửi thật. SMTP cần cấu hình từ nhà cung cấp; khởi động lại worker sau khi sửa. Không đặt mật khẩu, API key hoặc log trong public.'),
('p','Sao lưu SQL MySQL, thư mục ảnh và .env riêng. Trên hosting dùng web root public, HTTPS, production, tắt debug và mật khẩu thật. Không dùng artisan serve làm web server công khai. Cấu hình ngân hàng và AI do chủ đồ án tự điền theo lựa chọn đã thống nhất.')]),
('Kết quả kiểm thử',[
('p','Bộ kiểm thử đạt 81 ca với 779 assertions trên SQLite và MySQL 8.0.46. Bài kiểm thử MySQL chỉ dựng dữ liệu trong sole_store_test. SQL schema/demo được xuất từ migration và kiểm tra import riêng. Dữ liệu website được đối chiếu tồn, giữ chỗ, đơn và tiền.'),
('table',[['Nhóm kiểm tra','Điều kiện được xác minh'],['Đơn và kho','Không bán vượt tồn, giữ đúng SKU, hoàn tất/hủy không ghi lặp'],['Thanh toán','Đúng tổng, mã ngân hàng duy nhất, báo chuyển chưa thu, hoàn/tiền đến muộn'],['Hết hạn QR','Trả giữ/lượt mã một lần; không hủy đơn đã báo, thu, xác nhận hoặc đơn cũ'],['Tài khoản','Token riêng, băm, hết hạn, dùng một lần và thu hồi phiên cũ'],['Khách hàng','Đánh giá sau mua, duyệt/ẩn, escape; yêu cầu riêng tư và chống gửi lặp'],['Email và ảnh','Mail lỗi không mất đơn; outbox một lần; giữ ảnh cũ, đổi bìa và chặn ảnh khác sản phẩm']]),
('p','Đã kiểm tra trình duyệt: đăng nhập khách, gửi yêu cầu đổi size, admin phản hồi, gửi và duyệt đánh giá. Tài khoản/đơn minh họa có nhãn DEMO. Chưa thực hiện chuyển tiền ngân hàng, gửi SMTP thật hoặc gọi AI bằng khóa thật; kiểm thử không chứng minh các tích hợp này đã hoạt động ngoài môi trường đồ án.')]),
('Kết luận và hướng phát triển',[
('p','SOLE đã có luồng quản lý nhập–tồn–bán và kênh khách mua hàng. Bản cập nhật bổ sung quản lý đơn QR bỏ dở, phục hồi tài khoản, hàng đợi email, đánh giá đã mua và yêu cầu đổi/trả. Cơ sở dữ liệu và các quy tắc kho/tiền được kiểm tra bằng test tự động.'),
('p','Những việc cần dữ liệu hoặc dịch vụ thực tế: điền tài khoản ngân hàng, cấu hình SMTP, API AI nếu muốn sinh ngôn ngữ, ảnh đúng sản phẩm và chính sách cửa hàng được xác nhận. Vị trí giao hàng vẫn do cửa hàng cập nhật; chưa có GPS/API vận chuyển hoặc webhook ngân hàng. Chưa có nhiều kho, giá vốn/lợi nhuận, trả một phần hay khấu trừ chênh lệch đổi size tự động.'),
('h','Kịch bản bảo vệ khoảng 10 phút'),
('p','Giới thiệu mục tiêu và ba vai trò trong một phút. Trình bày khách lọc giày, giỏ, COD/QR và theo dõi trong hai phút. Minh họa nhập kho, giữ chỗ và hoàn tất trong hai phút. Trình bày hết hạn QR, quên mật khẩu, đánh giá và yêu cầu đổi/trả trong hai phút. Dành hai phút cho MySQL, test, báo cáo và một phút giải thích giới hạn.'),
('h','Tài liệu tham khảo và đối chiếu'),
('p','Tài liệu dự án: README.md; docs/PROJECT.md; DATABASE.md; TESTING.md; PAYMENT_DELIVERY.md; CUSTOMER_EXPERIENCE.md. Nguồn khung báo cáo: BaoCaoBTL/DE_CUONG_BTL_WEBSITE_QUAN_LY_BAN_HANG.docx. Tài liệu công nghệ: laravel.com/docs/12.x; dev.mysql.com/doc/refman/8.0/en; developer.mozilla.org. Không sao chép nhận xét giảng viên hoặc kết quả chấm điểm.')])]

doc=Document(); sec=doc.sections[0]
# Word's bundled template adds decorative borders to Title and Subtitle.
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=sec.bottom_margin=Inches(.7); sec.left_margin=sec.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name]; st.font.name='Arial'; st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_after=Pt(8)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.16
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(13)
doc.styles['Title'].font.size=Pt(24)
head=sec.header.paragraphs[0]; head.text='SOLE  |  ĐỒ ÁN CHUYÊN NGÀNH CÔNG NGHỆ PHẦN MỀM'; head.style='Caption'
head.runs[0].font.size=Pt(8); head.runs[0].font.color.rgb=RGBColor(0,0,0)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('SOLE  •  ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.add_paragraph('TRƯỜNG ĐẠI HỌC CÔNG NGHỆ ĐÔNG Á',style='Subtitle')
doc.add_paragraph('Báo cáo kỹ thuật website quản lý cửa hàng bán giày',style='Title')
doc.add_paragraph('PHP  JavaScript  MySQL',style='Subtitle')
doc.add_paragraph('Nhóm 6\nPhạm Quang Huy\nĐào Ngọc Tú')
doc.add_paragraph('Người hướng dẫn theo đề cương\nThS. Mai Văn Linh')
doc.add_paragraph('Bắc Ninh tháng 10 năm 2026')
doc.add_heading('Nội dung báo cáo',level=2)
doc.add_paragraph('Báo cáo mô tả phần mềm đã triển khai, thiết kế dữ liệu, quy tắc kho và thanh toán, các bổ sung cho khách hàng, kết quả kiểm thử và hướng dẫn chạy demo. Nội dung được cập nhật theo mã nguồn ngày 21/10/2026.')
doc.add_paragraph('Các tài khoản, chứng từ và nhận xét có nhãn DEMO chỉ dùng để minh họa. Thông tin ngân hàng, SMTP và API AI chưa được cấu hình để giao dịch thật.')

md=['# Báo cáo kỹ thuật website quản lý cửa hàng bán giày','\nNhóm 6 — Phạm Quang Huy và Đào Ngọc Tú. Cập nhật 21/10/2026.\n']
for title,blocks in pages:
    doc.add_page_break(); doc.add_heading(title,level=1); md.append('\n## '+title+'\n')
    for kind,value in blocks:
        if kind=='p': doc.add_paragraph(value); md.append(value+'\n')
        elif kind=='h': doc.add_heading(value,level=2); md.append('\n### '+value+'\n')
        elif kind=='img':
            if (OUT/value).exists():
                im=Image.open(OUT/value); maxw,maxh=6.8,3.55; w=min(maxw,maxh*im.width/im.height)
                doc.add_picture(str(OUT/value),width=Inches(w)); doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
                md.append('!['+value+']('+value+')\n')
        elif kind=='table':
            t=doc.add_table(rows=0,cols=len(value[0])); t.autofit=False
            for i,row in enumerate(value):
                cells=t.add_row().cells
                for j,text in enumerate(row):
                    cells[j].text=text
                    tcPr=cells[j]._tc.get_or_add_tcPr()
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'324E3F' if i==0 else ('F3F5F2' if i%2 else 'FFFFFF')); tcPr.append(shade)
                    borders=OxmlElement('w:tcBorders')
                    for side in ['top','left','bottom','right']:
                        edge=OxmlElement('w:'+side); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'4'); edge.set(qn('w:color'),'D9D9D9'); borders.append(edge)
                    tcPr.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        edge=OxmlElement('w:'+side); edge.set(qn('w:w'),'95'); edge.set(qn('w:type'),'dxa'); margins.append(edge)
                    tcPr.append(margins)
                    for p in cells[j].paragraphs:
                        p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.08
                        for r in p.runs:
                            r.font.size=Pt(10); r.font.bold=i==0; r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '222222')
                trPr=t.rows[-1]._tr.get_or_add_trPr(); noSplit=OxmlElement('w:cantSplit'); trPr.append(noSplit)
                if i==0: trPr.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph('')
            md.extend(['| '+' | '.join(row)+' |' for row in value])
    
doc.save(OUT/'BAO_CAO_KY_THUAT_SOLE.docx')
(OUT/'BAO_CAO_KY_THUAT_SOLE.md').write_text('\n'.join(md),encoding='utf-8')
(BUILD/'report-pages.json').write_text(json.dumps(pages,ensure_ascii=False),encoding='utf-8')
print('Created report and diagrams:',OUT)
